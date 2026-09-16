"""report.py — Geração do relatório Excel ao final de cada execução.

Cria um arquivo .xlsx na pasta de Downloads contendo, para cada chave:
chave de acesso, número da nota fiscal, status e mensagem de retorno.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import config

_CABECALHO = ["Chave de acesso", "Número da nota fiscal", "Status", "Mensagem"]
_FILL_TITULO = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")


def _estilo_titulo(plan) -> None:
    """Aplica formatação do cabeçalho e ajusta as larguras das colunas."""
    for celula in plan[1]:
        celula.font = Font(bold=True, color="FFFFFF", size=11)
        celula.fill = _FILL_TITULO
        celula.alignment = Alignment(horizontal="center", vertical="center")

    larguras = [46, 14, 12, 60]
    for indice, largura in enumerate(larguras, start=1):
        plan.column_dimensions[get_column_letter(indice)].width = largura

    plan.freeze_panes = "A2"


def gerar_relatorio(resultados: list[dict], destino: Path) -> Path:
    """Gera o relatório Excel com os resultados da execução.

    Args:
        resultados: lista de dicts com chaves:
            chave, numero, status (sucesso/falha), mensagem.
        destino: pasta onde o arquivo deve ser salvo.

    Returns:
        Path do arquivo .xlsx gerado (nome com carimbo de data/hora).
    """
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    caminho = destino / f"{config.PREFIXO_RELATORIO}_{timestamp}.xlsx"

    workbook = Workbook()
    plan = workbook.active
    plan.title = "Resultado"
    plan.append(_CABECALHO)

    for resultado in resultados:
        plan.append(
            [
                resultado.get("chave", ""),
                resultado.get("numero", ""),
                resultado.get("status", ""),
                resultado.get("mensagem", ""),
            ]
        )

    _estilo_titulo(plan)
    workbook.save(caminho)
    return caminho